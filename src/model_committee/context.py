"""Context preparation: load user-supplied files as text, cap their size, flag likely prompt injection."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from pathlib import Path

MAX_TOTAL_CHARS = 120_000
TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".yaml", ".yml", ".xml", ".html", ".htm",
                 ".py", ".js", ".ts", ".sql", ".log", ".rst", ".ini", ".toml"}

INJECTION_PATTERNS = [
    r"ignore (all |any )?(the )?(previous|prior|above) (instructions|prompts?)",
    r"disregard (all |any )?(the )?(previous|prior|above|system)",
    r"you are now (a|an|the) ",
    r"new (system )?instructions?:",
    r"(reveal|print|show) (your|the) (system )?prompt",
    r"<\|im_start\|>|<\|system\|>|\[INST\]",
    r"do not (tell|inform) the user",
    r"send (this|the|all) .{0,40}(to|at) [\w.+-]+@[\w-]+\.",
]
_INJECTION = [re.compile(p, re.IGNORECASE) for p in INJECTION_PATTERNS]


class ContextError(ValueError):
    pass


@dataclass
class ContextFile:
    name: str
    source_path: str
    text: str
    chars: int
    truncated: bool = False
    injection_flags: list[str] = field(default_factory=list)


@dataclass
class ContextBundle:
    files: list[ContextFile] = field(default_factory=list)

    @property
    def names(self) -> list[str]:
        return [f.name for f in self.files]

    @property
    def flagged(self) -> list[ContextFile]:
        return [f for f in self.files if f.injection_flags]


def scan_injection(text: str) -> list[str]:
    return sorted({m.group(0)[:80] for rx in _INJECTION for m in rx.finditer(text)})


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n\n".join(f"[page {i + 1}]\n{page.extract_text() or ''}" for i, page in enumerate(reader.pages))


def _read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    out = io.StringIO()
    for ws in wb.worksheets:
        out.write(f"[sheet {ws.title}]\n")
        writer = csv.writer(out)
        for row in ws.iter_rows(values_only=True):
            writer.writerow(["" if v is None else v for v in row])
    return out.getvalue()


def read_file_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix in {".xlsx", ".xlsm"}:
        return _read_xlsx(path)
    if suffix in TEXT_SUFFIXES or not suffix:
        return path.read_text(encoding="utf-8", errors="replace")
    raise ContextError(f"unsupported context file type: {path.name} (supported: pdf, xlsx, and text formats)")


def load_context(paths: list[str], *, max_total_chars: int = MAX_TOTAL_CHARS) -> ContextBundle:
    """Every member receives the same bundle; the run records exactly what each one got."""
    bundle = ContextBundle()
    if not paths:
        return bundle
    per_file = max_total_chars // len(paths)
    seen: set[str] = set()
    for raw in paths:
        path = Path(raw).expanduser()
        if not path.is_file():
            raise ContextError(f"context file not found: {raw}")
        text = read_file_text(path)
        name = path.name
        while name in seen:
            name = f"_{name}"
        seen.add(name)
        truncated = len(text) > per_file
        bundle.files.append(
            ContextFile(
                name=name,
                source_path=str(path.resolve()),
                text=text[:per_file],
                chars=len(text),
                truncated=truncated,
                injection_flags=scan_injection(text),
            )
        )
    return bundle
