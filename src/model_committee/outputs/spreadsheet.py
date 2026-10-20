"""xlsx workbook: claim matrix, evidence, members, decision matrix and validation, one sheet each."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from model_committee.outputs import (
    CLASSIFICATION_ORDER,
    STATUS_LABEL,
    effective_classification,
    member_names,
    member_order,
    safe_url,
    total_tokens,
)
from model_committee.schemas import MemberReport, MemberStatus, Synthesis, ValidationReport

_MAX_CELL = 32_000  # Excel's hard limit is 32,767 characters
_HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_FILLS = {
    "agreement": "E3F1E7",
    "disagreement": "FBEADB",
    "unique": "E2ECF8",
    "contested": "ECECEC",
    "support": "E3F1E7",
    "reject": "F8D9D9",
    "uncertain": "FFF1CC",
    "untraceable": "F8D9D9",
    "warning": "FFF1CC",
}
_WRAP_TOP = Alignment(wrap_text=True, vertical="top")


def _clean(value: object) -> object:
    """Strip characters Excel rejects and cap length. Non-strings pass through unchanged."""
    if not isinstance(value, str):
        return value
    value = ILLEGAL_CHARACTERS_RE.sub("", value)
    return value[: _MAX_CELL - 1] + "…" if len(value) > _MAX_CELL else value


def _append(ws: Worksheet, values: list[object]) -> int:
    """Append a row of untrusted values. Strings are always stored as text, never as formulas."""
    ws.append([_clean(v) for v in values])
    row = ws.max_row
    for cell in ws[row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"  # a leading '=' from model output must not become a live formula
            cell.alignment = _WRAP_TOP
        else:
            cell.alignment = Alignment(vertical="top")
    return row


def _new_sheet(wb: Workbook, title: str, header: list[str]) -> Worksheet:
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    return ws


def _finish(ws: Worksheet, widths: dict[int, float] | None = None, max_width: float = 60) -> None:
    """Autofilter over the used range and column widths sized to content (overridable per column)."""
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{max(ws.max_row, 1)}"
    for idx, column in enumerate(ws.iter_cols(min_row=1, max_row=ws.max_row), start=1):
        longest = max((len(str(c.value)) for c in column if c.value is not None), default=8)
        width = (widths or {}).get(idx) or min(max(longest + 2, 10), max_width)
        ws.column_dimensions[get_column_letter(idx)].width = width


def _class_rule(ws: Worksheet, col: int, values: list[str]) -> None:
    """Conditional fill on a text column, so colours follow the value when rows are sorted or edited."""
    if ws.max_row < 2:
        return
    letter = get_column_letter(col)
    rng = f"{letter}2:{letter}{ws.max_row}"
    for value in values:
        fill = PatternFill("solid", start_color=_FILLS[value], end_color=_FILLS[value])
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LOWER(${letter}2)="{value}"'], fill=fill))


# ---------------------------------------------------------------- sheets


def _claims_sheet(
    wb: Workbook, synthesis: Synthesis, keys: list[str], names: dict[str, str], statuses: dict[str, MemberStatus]
) -> None:
    header = (
        ["Claim ID", "Statement", "Classification"]
        + [f"{names[k]} stance" for k in keys]
        + [
            "Support",
            "Reject",
            "Uncertain",
            "Model consensus",
            "Independent sources",
            "Source domains",
            "Disagreement type",
            "Disagreement cause",
        ]
    )
    ws = _new_sheet(wb, "Claims", header)
    for c in synthesis.matrix.claims:
        stances = []
        for k in keys:
            pos = c.positions.get(k)
            if pos is not None:
                stances.append(pos.stance)
            elif statuses.get(k, MemberStatus.succeeded) is not MemberStatus.succeeded:
                stances.append(f"({statuses[k].value})")
            else:
                stances.append("not_mentioned")
        row = _append(
            ws,
            [
                c.id,
                c.statement,
                effective_classification(c).value,
                *stances,
                c.support,
                c.reject,
                c.uncertain,
                c.model_consensus,
                c.independent_sources,
                ", ".join(c.source_domains),
                c.disagreement_type.value if c.disagreement_type else None,
                c.disagreement_cause,
            ],
        )
        ws.cell(row, 4 + len(keys) + 3).number_format = "0%"
    _class_rule(ws, 3, [c.value for c in CLASSIFICATION_ORDER])
    for i in range(len(keys)):
        _class_rule(ws, 4 + i, ["support", "reject", "uncertain"])
    widths = {1: 10, 2: 60, 3: 15} | {4 + i: 16 for i in range(len(keys))}
    _finish(ws, widths)


def _claim_index(synthesis: Synthesis) -> dict[str, list[str]]:
    """'<member_key>/<claim_id>' -> canonical claim ids that reference it."""
    index: dict[str, list[str]] = {}
    for c in synthesis.matrix.claims:
        for pos in c.positions.values():
            for ref in pos.claim_refs:
                index.setdefault(ref, []).append(c.id)
    return index


def _evidence_sheet(wb: Workbook, synthesis: Synthesis, members: list[MemberReport], names: dict[str, str]) -> None:
    ws = _new_sheet(
        wb,
        "Evidence",
        ["Member", "Canonical claim", "Member claim", "Member claim statement", "Source", "URL", "Quote"],
    )
    index = _claim_index(synthesis)
    for m in members:
        for mc in m.findings.claims if m.findings else []:
            canonical = ", ".join(index.get(f"{m.key}/{mc.id}", []))
            for e in mc.evidence:
                row = _append(ws, [names[m.key], canonical or None, mc.id, mc.statement, e.source, e.url, e.quote])
                url = safe_url(e.url)
                if url:
                    ws.cell(row, 6).hyperlink = url
                    ws.cell(row, 6).style = "Hyperlink"
                    ws.cell(row, 6).alignment = _WRAP_TOP
    _finish(ws, {2: 12, 3: 10, 4: 45, 5: 32, 6: 45, 7: 60})


def _members_sheet(wb: Workbook, members: list[MemberReport], names: dict[str, str]) -> None:
    ws = _new_sheet(
        wb,
        "Members",
        [
            "Key",
            "Label",
            "Spec",
            "Vendor",
            "Status",
            "Attempts",
            "Duration (s)",
            "Cost (USD)",
            "Tokens",
            "Web searches",
            "Confidence",
            "Recommendation",
            "Summary",
            "Error",
        ],
    )
    for m in members:
        t = m.telemetry
        f = m.findings
        row = _append(
            ws,
            [
                m.key,
                names[m.key],
                m.spec,
                m.vendor,
                STATUS_LABEL[m.status],
                m.attempts,
                round(t.duration_s, 1) if t else None,
                t.cost_usd if t else None,
                total_tokens(m),
                t.web_searches if t else None,
                f.confidence if f else None,
                f.recommendation if f else None,
                f.summary if f else None,
                m.error or m.parse_error,
            ],
        )
        ws.cell(row, 8).number_format = '"$"#,##0.00'
        ws.cell(row, 9).number_format = "#,##0"
        ws.cell(row, 11).number_format = "0%"
        if m.status is not MemberStatus.succeeded:
            ws.cell(row, 5).fill = PatternFill("solid", fgColor=_FILLS["reject"])
    _finish(ws, {12: 50, 13: 50, 14: 50})


def _decision_sheet(wb: Workbook, synthesis: Synthesis) -> None:
    rows = synthesis.narrative.decision_matrix
    criteria: list[str] = []
    for r in rows:
        criteria += [k for k in r.scores if k not in criteria]
    ws = _new_sheet(wb, "Decision Matrix", ["Option", *criteria, "Verdict"])
    for r in rows:
        _append(ws, [r.option, *(r.scores.get(k) for k in criteria), r.verdict])
    _finish(ws, max_width=40)


def _validation_sheet(wb: Workbook, validation: ValidationReport, names: dict[str, str]) -> None:
    ws = _new_sheet(
        wb,
        "Validation",
        ["Type", "URL / message", "Traced to members", "Fetched", "HTTP status", "Quote found", "Note"],
    )
    for url in validation.untraceable_citations:
        _append(ws, ["untraceable", url, None, None, None, None, "Cited by the synthesis but by no member."])
    if not validation.unique_findings_preserved:
        _append(ws, ["warning", "A unique member finding is missing from the synthesis.", None, None, None, None, None])
    for w in validation.warnings:
        _append(ws, ["warning", w, None, None, None, None, None])
    for c in validation.citations:
        traced = ", ".join(names.get(k, k) for k in c.traced_to_members)
        _append(ws, ["citation", c.url, traced or "(none)", c.fetched, c.http_status, c.quote_found, c.note])
    _class_rule(ws, 1, ["untraceable", "warning"])
    _finish(ws, {2: 60, 7: 45})


def write_xlsx(path: Path, *, synthesis: Synthesis, members: list[MemberReport], validation: ValidationReport) -> None:
    keys = member_order(synthesis, members)
    names = member_names(members, keys)
    statuses = {m.key: m.status for m in members}
    wb = Workbook()
    wb.remove(wb.active)
    _claims_sheet(wb, synthesis, keys, names, statuses)
    _evidence_sheet(wb, synthesis, members, names)
    _members_sheet(wb, members, names)
    _decision_sheet(wb, synthesis)
    _validation_sheet(wb, validation, names)
    wb.properties.title = synthesis.narrative.headline[:250]
    wb.properties.subject = synthesis.objective[:250]
    wb.save(Path(path))
