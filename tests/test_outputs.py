"""Output skills: HTML report, xlsx workbook, pptx deck."""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

from model_committee.outputs import render_outputs
from model_committee.outputs.deck import ROWS_PER_TABLE, write_pptx
from model_committee.outputs.report import md, render_html
from model_committee.outputs.spreadsheet import write_xlsx
from model_committee.schemas import DecisionRow, Deliverable, MemberStatus
from tests.sample_data import M1, M3, UNTRACEABLE_URL, make_sample

XSS = "<script>alert('xss')</script>"

VARIANTS = {
    "full": {},
    "failed": {"failed_member": True},
    "empty": {"empty": True},
}


@pytest.fixture(params=list(VARIANTS))
def sample(request):
    return make_sample(**VARIANTS[request.param])


def _inject(sample: dict) -> dict:
    """Put a script payload into every kind of untrusted text the renderers show."""
    members = sample["members"]
    m1 = next(m for m in members if m.key == M1)
    m1.report_markdown += f"\n\n{XSS}\n\n[click](javascript:alert(1)) ![pixel](https://tracker.example.com/p.gif)\n"
    m1.findings.summary = XSS
    synthesis = sample["synthesis"]
    if synthesis.matrix.claims:
        synthesis.matrix.claims[0].statement = f"Cheapest {XSS}"
    synthesis.narrative.answer += XSS
    return sample


# ---------------------------------------------------------------- render_outputs


@pytest.mark.parametrize(
    ("deliverable", "expected"),
    [
        (Deliverable.none, []),
        (Deliverable.report, [("report.html", "report")]),
        (Deliverable.spreadsheet, [("committee.xlsx", "spreadsheet")]),
        (Deliverable.deck, [("committee-deck.pptx", "deck")]),
        (
            Deliverable.all,
            [("report.html", "report"), ("committee.xlsx", "spreadsheet"), ("committee-deck.pptx", "deck")],
        ),
    ],
)
def test_render_outputs_writes_requested_files(tmp_path: Path, deliverable, expected):
    out = tmp_path / "out"
    written = render_outputs(deliverable, out_dir=out, **make_sample())
    assert written == expected
    for name, _ in written:
        assert (out / name).stat().st_size > 0


# ---------------------------------------------------------------- HTML


def test_html_has_every_section(sample):
    page = render_html(**sample)
    for sid in [
        "answer",
        "agreement",
        "disagreement",
        "unique",
        "contested",
        "consensus",
        "decision",
        "risks",
        "validation",
        "members",
    ]:
        assert f'<section id="{sid}">' in page
    for m in sample["members"]:
        assert f'id="member-{m.key}"' in page


def test_html_content_for_full_sample():
    sample = make_sample()
    page = render_html(**sample)
    # disagreement cause, type and recommended handling
    assert "Give ranges, drivers and leading indicators." in page  # forecast handling
    assert "Expose the assumptions and rerun the scenarios." in page  # assumption handling
    assert "council record" in page
    # unique finding attributed to its member
    assert "Raised only by <strong>Codex (CLI)</strong>" in page
    # consensus vs evidence explains the difference and flags thin evidence
    assert "<strong>not</strong> independent confirmation" in page
    assert "Thin evidence" in page
    # untraceable citation prominently flagged
    assert "Untraceable citations" in page and 'role="alert"' in page
    assert UNTRACEABLE_URL in page
    # decision matrix, risks, actions, telemetry, full report rendered from markdown
    assert "Not recommended" in page
    assert "Rebuild the payback model" in page
    assert "$1.84" in page and "6m 53s" in page
    assert "<table>" in page and "<h2>Cost comparison</h2>" in page


def test_html_failed_member_shown_not_dropped():
    sample = make_sample(failed_member=True)
    page = render_html(**sample)
    failed = next(m for m in sample["members"] if m.key == M3)
    assert failed.status is MemberStatus.failed
    assert 'class="member failed"' in page
    assert "quota exhausted" in page
    assert "2 of 3 succeeded" in page


def test_html_empty_sample_says_so():
    page = render_html(**make_sample(empty=True))
    assert "No decision matrix was produced." in page
    assert "No claims to compare." in page
    assert "No recommendation made." in page


def test_html_escapes_untrusted_text():
    page = render_html(**_inject(make_sample()))
    assert "<script" not in page.lower()
    assert "&lt;script&gt;alert(&#x27;xss&#x27;)&lt;/script&gt;" in page
    assert "javascript:" not in page
    assert "<img" not in page


def test_html_has_no_external_resources(sample):
    page = render_html(**_inject(sample))
    # Only <a href> may point off-page; nothing may be loaded.
    assert not re.search(r"""\bsrc\s*=""", page, re.I)
    assert not re.search(r"<link\b", page, re.I)
    assert not re.search(r"@import|url\(", page, re.I)
    for href in re.findall(r'href="([^"]*)"', page):
        assert href.startswith(("#", "http://", "https://", "mailto:")), href


def test_html_is_print_and_dark_aware():
    page = render_html(**make_sample())
    assert "@media print" in page
    assert "prefers-color-scheme:dark" in page
    assert 'name="viewport"' in page


@pytest.mark.parametrize(
    ("source", "forbidden"),
    [
        ('<img src=x onerror="alert(1)">', "<img"),
        ("[x](javascript:alert(1))", "javascript:"),
        ("[x](&#106;avascript:alert(1))", "avascript:alert"),
        ('<div onclick="steal()">hi</div>', "<div"),
        ("<iframe src='https://evil.example.com'></iframe>", "<iframe"),
        ("<style>body{display:none}</style>", "<style"),
    ],
)
def test_markdown_sanitizer_blocks_markup(source, forbidden):
    rendered = md(source)
    assert forbidden not in rendered.replace("&lt;", "")  # escaped text is fine, live markup is not


def test_markdown_keeps_safe_formatting():
    rendered = md("# Title\n\n**bold** [site](https://example.com)\n\n```\na < b\n```")
    assert "<h1>Title</h1>" in rendered
    assert "<strong>bold</strong>" in rendered
    assert 'href="https://example.com"' in rendered
    assert "a &lt; b" in rendered


# ---------------------------------------------------------------- XLSX


def _evidence_count(sample: dict) -> int:
    return sum(len(c.evidence) for m in sample["members"] if m.findings for c in m.findings.claims)


def test_xlsx_structure(tmp_path: Path, sample):
    path = tmp_path / "c.xlsx"
    write_xlsx(path, synthesis=sample["synthesis"], members=sample["members"], validation=sample["validation"])
    wb = load_workbook(path)
    assert wb.sheetnames == ["Claims", "Evidence", "Members", "Decision Matrix", "Validation"]
    claims = sample["synthesis"].matrix.claims
    assert wb["Claims"].max_row == len(claims) + 1
    assert wb["Evidence"].max_row == _evidence_count(sample) + 1
    assert wb["Members"].max_row == len(sample["members"]) + 1
    assert wb["Decision Matrix"].max_row == len(sample["synthesis"].narrative.decision_matrix) + 1
    for ws in wb.worksheets:
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref


def test_xlsx_claim_columns(tmp_path: Path):
    sample = make_sample(failed_member=True)
    path = tmp_path / "c.xlsx"
    write_xlsx(path, synthesis=sample["synthesis"], members=sample["members"], validation=sample["validation"])
    ws = load_workbook(path)["Claims"]
    header = [c.value for c in ws[1]]
    for m in sample["members"]:
        assert f"{m.label} stance" in header
    rows = {r[0]: dict(zip(header, r, strict=True)) for r in ws.iter_rows(min_row=2, values_only=True)}
    k3 = rows["K3"]
    assert k3["Classification"] == "disagreement"
    assert k3["Disagreement type"] == "forecast"
    assert k3["Gemini 3.8 Flash (agy) stance"] == "(failed)"
    assert k3["Support"] == 1 and k3["Reject"] == 1
    assert ws.conditional_formatting  # classification fill


def test_xlsx_never_writes_formulas_from_model_text(tmp_path: Path):
    sample = make_sample()
    sample["synthesis"].matrix.claims[0].statement = '=HYPERLINK("https://evil.example.com","x")'
    path = tmp_path / "c.xlsx"
    write_xlsx(path, synthesis=sample["synthesis"], members=sample["members"], validation=sample["validation"])
    cell = load_workbook(path)["Claims"]["B2"]
    assert cell.data_type == "s"
    assert cell.value.startswith("=HYPERLINK")


def test_xlsx_validation_lists_untraceable(tmp_path: Path):
    sample = make_sample()
    path = tmp_path / "c.xlsx"
    write_xlsx(path, synthesis=sample["synthesis"], members=sample["members"], validation=sample["validation"])
    rows = list(load_workbook(path)["Validation"].iter_rows(min_row=2, values_only=True))
    assert ("untraceable", UNTRACEABLE_URL) in [(r[0], r[1]) for r in rows]


# ---------------------------------------------------------------- PPTX


def _write_deck(tmp_path: Path, sample: dict) -> Presentation:
    path = tmp_path / "d.pptx"
    write_pptx(
        path,
        synthesis=sample["synthesis"],
        members=sample["members"],
        validation=sample["validation"],
        config=sample["config"],
    )
    return Presentation(str(path))


def _slide_text(slide) -> str:
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            parts.append(shape.text_frame.text)
        if shape.has_table:
            parts += [c.text for row in shape.table.rows for c in row.cells]
    return "\n".join(parts)


def test_pptx_structure(tmp_path: Path, sample):
    prs = _write_deck(tmp_path, sample)
    assert len(prs.slides) > 6
    assert abs(prs.slide_width / prs.slide_height - 16 / 9) < 0.01
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:  # text boxes are only created for real content
                assert shape.text_frame.text.strip(), f"empty text box on slide {slide.slide_id}"
            # nothing placed outside the slide
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= prs.slide_width + Emu(10)
            assert shape.top + shape.height <= prs.slide_height + Emu(10)
    text = "\n".join(_slide_text(s) for s in prs.slides)
    for m in sample["members"]:
        assert f"Appendix: {m.label}" in text


def test_pptx_full_content(tmp_path: Path):
    text = "\n".join(_slide_text(s) for s in _write_deck(tmp_path, make_sample()).slides)
    for needle in [
        "Choose Contoso Freight",
        "Agreement",
        "Disagreements",
        "Handling: Give ranges",
        "Raised only by Codex (CLI)",
        "Consensus vs evidence",
        "Decision matrix",
        "Unresolved risks",
        "not traceable",
    ]:
        assert needle in text


def test_pptx_failed_member_slide(tmp_path: Path):
    prs = _write_deck(tmp_path, make_sample(failed_member=True))
    last = _slide_text(prs.slides[-1])
    assert "Failed" in last and "quota exhausted" in last


def test_pptx_splits_long_tables(tmp_path: Path):
    sample = make_sample()
    sample["synthesis"].narrative.decision_matrix = [
        DecisionRow(option=f"Option {i}", scores={"Cost": f"${i}", "Risk": "Medium " * 30}, verdict="Maybe")
        for i in range(20)
    ]
    prs = _write_deck(tmp_path, sample)
    tables = [
        shape.table
        for slide in prs.slides
        if _slide_text(slide).startswith("Decision matrix")
        for shape in slide.shapes
        if shape.has_table
    ]
    assert len(tables) == 3  # 8 + 8 + 4
    assert all(len(t.rows) - 1 <= ROWS_PER_TABLE for t in tables)
    assert sum(len(t.rows) - 1 for t in tables) == 20
    long_cells = [t.cell(1, 2).text for t in tables]
    assert all(c.endswith("…") for c in long_cells)  # over-long cell text truncated, not overflowing


def test_pptx_paginates_long_lists(tmp_path: Path):
    sample = make_sample()
    claim = sample["synthesis"].matrix.claims[0]
    sample["synthesis"].matrix.claims = [
        claim.model_copy(update={"id": f"A{i}", "statement": f"Agreed claim number {i}. " * 6}) for i in range(15)
    ]
    sample["synthesis"].narrative.unresolved_risks = [f"Risk {i}: " + "detail " * 25 for i in range(14)]
    prs = _write_deck(tmp_path, sample)
    titles = [_slide_text(s).split("\n")[0] for s in prs.slides]
    assert titles.count("Agreement") == 1 and "Agreement (continued)" in titles
    assert "Risks and next actions (continued)" in titles
    text = "\n".join(_slide_text(s) for s in prs.slides)
    for i in range(15):
        assert f"A{i} " in text  # every claim kept, none silently dropped


def test_answer_markdown_is_rendered_not_shown_raw():
    from model_committee.outputs import plain
    from model_committee.outputs.report import render_html

    data = make_sample()
    data["synthesis"].narrative.answer = "**Default: pgvector.** First para.\n\n- point one\n- point two"
    page = render_html(**data)
    assert "<strong>Default: pgvector.</strong>" in page and "<li>point one</li>" in page
    assert "**Default" not in page
    assert plain("**bold** and *it* with [link](https://x.example) `code`") == "bold and it with link code"
